#!/usr/bin/env python3
"""
Generuje docs/output/onboarding-windows.docx — instrukcja startowa dla Windows.

Uruchomienie:
    python3 system/scripts/generate-onboarding-windows-docx.py

Wymagania (instalowane automatycznie):
    python-docx, pillow
"""

import subprocess
import sys
import os
import tempfile

# ── Auto-instalacja zależności ────────────────────────────────────────────────

def pip_install(package):
    print(f"Instaluję {package}...")
    subprocess.check_call([sys.executable, "-m", "pip", "install", package],
                          stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)

try:
    from docx import Document
except ImportError:
    pip_install("python-docx")
    from docx import Document

try:
    from PIL import Image, ImageDraw, ImageFont
except ImportError:
    pip_install("pillow")
    from PIL import Image, ImageDraw, ImageFont

from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

IMG_WIDTH_PX = 520
IMG_WIDTH_IN  = 5.0

# ── Helpers: czcionki ─────────────────────────────────────────────────────────

def get_font(size=13):
    candidates = [
        '/System/Library/Fonts/Helvetica.ttc',
        '/System/Library/Fonts/SFNSText.ttf',
        '/Library/Fonts/Arial.ttf',
    ]
    for path in candidates:
        if os.path.exists(path):
            try:
                return ImageFont.truetype(path, size)
            except Exception:
                pass
    return ImageFont.load_default()

def get_mono_font(size=13):
    candidates = [
        '/System/Library/Fonts/Menlo.ttc',
        '/System/Library/Fonts/Monaco.ttf',
        '/Library/Fonts/Courier New.ttf',
    ]
    for path in candidates:
        if os.path.exists(path):
            try:
                return ImageFont.truetype(path, size)
            except Exception:
                pass
    return ImageFont.load_default()

# ── Helpers: mockupy ─────────────────────────────────────────────────────────

def create_powershell_mockup(lines, width=IMG_WIDTH_PX):
    """Tworzy mockup okna PowerShell (ciemne tło, niebieskie akcenty)."""
    try:
        mono = get_mono_font(13)
        font = get_font(12)

        line_h = 22
        pad = 14
        title_h = 30
        height = title_h + pad + len(lines) * line_h + pad + 10

        img = Image.new('RGB', (width, height), (12, 12, 12))
        draw = ImageDraw.Draw(img)

        # Pasek tytułu — niebieski jak PowerShell
        draw.rectangle([(0, 0), (width, title_h)], fill=(1, 36, 86))
        draw.text((12, 8), "Windows PowerShell", font=font, fill=(220, 220, 220))
        # Przyciski zamknięcia
        draw.rectangle([(width - 42, 6), (width - 8, 24)], fill=(196, 43, 28))
        draw.text((width - 33, 8), "X", font=font, fill=(255, 255, 255))

        y = title_h + pad
        for line in lines:
            if line.startswith('>>>'):
                draw.text((pad, y), line[3:].strip(), font=mono, fill=(130, 130, 130))
            elif line.startswith('#'):
                draw.text((pad, y), line, font=mono, fill=(80, 80, 80))
            else:
                draw.text((pad, y), f"PS C:\\Users\\Ty> {line}", font=mono, fill=(220, 220, 220))
            y += line_h

        draw.rectangle([(pad, y), (pad + 9, y + 15)], fill=(200, 200, 200))

        out = tempfile.mktemp(suffix='.png')
        img.save(out, 'PNG')
        return out
    except Exception as e:
        print(f"  Mockup PowerShell nie powiódł się: {e}")
        return None


def create_windsurf_windows_mockup(width=IMG_WIDTH_PX):
    """Mockup Windsurf na Windows — pokazuje gdzie jest terminal."""
    try:
        font = get_font(12)
        font_sm = get_font(10)
        font_bold = get_font(14)

        height = 300
        editor_split = 175

        img = Image.new('RGB', (width, height), (30, 30, 30))
        draw = ImageDraw.Draw(img)

        # ── Pasek tytułu Windows ──
        draw.rectangle([(0, 0), (width, 28)], fill=(32, 32, 32))
        draw.text((8, 6), "Windsurf", font=font, fill=(200, 200, 200))
        # Przyciski Windows (minimize, maximize, close)
        for i, (label, color) in enumerate([('-', (80,80,80)), ('□', (80,80,80)), ('X', (196,43,28))]):
            bx = width - 105 + i * 35
            draw.rectangle([(bx, 4), (bx + 30, 24)], fill=color)
            draw.text((bx + 10, 6), label, font=font, fill=(220, 220, 220))

        # ── Menu bar ──
        draw.rectangle([(0, 28), (width, 52)], fill=(48, 48, 48))
        items = [('File', False), ('Edit', False), ('View', False),
                 ('Terminal', True), ('Help', False)]
        x = 8
        for label, highlight in items:
            color = (255, 215, 0) if highlight else (200, 200, 200)
            draw.text((x, 34), label, font=font, fill=color)
            x += len(label) * 7 + 18

        # ── Edytor ──
        draw.rectangle([(0, 52), (width, editor_split)], fill=(35, 35, 35))
        draw.text((16, 70), "# Twój kod będzie tutaj", font=font, fill=(70, 70, 70))
        draw.text((16, 95), "def hello():", font=font, fill=(80, 80, 80))
        draw.text((32, 117), 'print("Hello!")', font=font, fill=(80, 80, 80))

        # ── Separator ──
        draw.rectangle([(0, editor_split), (width, editor_split + 2)], fill=(80, 80, 80))

        # ── Panel terminala ──
        draw.rectangle([(0, editor_split + 2), (width, editor_split + 22)],
                        fill=(42, 42, 42))
        draw.text((10, editor_split + 5), "TERMINAL   PROBLEMS   OUTPUT",
                  font=font_sm, fill=(160, 160, 160))
        draw.rectangle([(10, editor_split + 20), (72, editor_split + 22)],
                        fill=(80, 180, 255))

        # ── Zawartość terminala (PowerShell) ──
        draw.rectangle([(0, editor_split + 22), (width, height)], fill=(12, 12, 12))
        draw.text((10, editor_split + 30), "PS C:\\Users\\Ty\\Onboarding>",
                  font=get_mono_font(12), fill=(80, 180, 255))
        draw.rectangle([(10, editor_split + 54), (18, editor_split + 68)],
                        fill=(200, 200, 200))

        # ── Strzałka ──
        ax = width - 20
        ay = editor_split + 55
        draw.polygon([(ax - 30, ay), (ax, ay - 14), (ax, ay + 14)], fill=(255, 60, 60))
        draw.rectangle([(ax - 100, ay - 7), (ax - 30, ay + 7)], fill=(255, 60, 60))
        draw.text((ax - 190, ay - 26), "Tu wpisuj", font=font_bold, fill=(255, 60, 60))
        draw.text((ax - 190, ay - 8), "komendy!", font=font, fill=(255, 60, 60))

        # ── Legenda ──
        draw.text((8, height - 38), "Jak otworzyć terminal:", font=font_sm,
                  fill=(180, 180, 180))
        draw.text((8, height - 22), "Terminal  →  New Terminal   (lub Ctrl + `)",
                  font=font_sm, fill=(255, 215, 0))

        out = tempfile.mktemp(suffix='.png')
        img.save(out, 'PNG')
        return out
    except Exception as e:
        print(f"  Mockup Windsurf Windows nie powiódł się: {e}")
        return None


# ── Helpers: dokument ────────────────────────────────────────────────────────

def set_font(run, name='Calibri', size=11, bold=False, italic=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_formcheckbox(paragraph):
    run = paragraph.add_run()
    fldChar_begin = OxmlElement('w:fldChar')
    fldChar_begin.set(qn('w:fldCharType'), 'begin')
    ffData = OxmlElement('w:ffData')
    for tag, attrs in [
        ('w:name', {'w:val': ''}),
        ('w:enabled', {}),
        ('w:calcOnExit', {'w:val': '0'}),
    ]:
        el = OxmlElement(tag)
        for k, v in attrs.items():
            el.set(qn(k), v)
        ffData.append(el)
    checkBox = OxmlElement('w:checkBox')
    for tag, attrs in [
        ('w:sizeAuto', {}),
        ('w:default', {'w:val': '0'}),
        ('w:checked', {'w:val': '0'}),
    ]:
        el = OxmlElement(tag)
        for k, v in attrs.items():
            el.set(qn(k), v)
        checkBox.append(el)
    ffData.append(checkBox)
    fldChar_begin.append(ffData)
    run._r.append(fldChar_begin)
    run2 = paragraph.add_run()
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = ' FORMCHECKBOX '
    run2._r.append(instrText)
    for ftype in ('separate', 'end'):
        r = paragraph.add_run()
        fc = OxmlElement('w:fldChar')
        fc.set(qn('w:fldCharType'), ftype)
        r._r.append(fc)


def add_section_heading(doc, text):
    doc.add_paragraph()
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    r = h.add_run(text)
    set_font(r, bold=True, size=14, color=(31, 73, 125))
    pPr = h._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    b = OxmlElement('w:bottom')
    b.set(qn('w:val'), 'single')
    b.set(qn('w:sz'), '6')
    b.set(qn('w:space'), '1')
    b.set(qn('w:color'), '1F497D')
    pBdr.append(b)
    pPr.append(pBdr)


def add_step(doc, number, title, description, sub_steps=None,
             image_path=None, image_caption=None):
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(12)
    para.paragraph_format.space_after = Pt(3)
    add_formcheckbox(para)
    para.add_run('  ')
    title_run = para.add_run(f'Krok {number} — {title}')
    set_font(title_run, bold=True, size=12)

    if description:
        d = doc.add_paragraph(description)
        d.paragraph_format.left_indent = Cm(1.2)
        d.paragraph_format.space_before = Pt(2)
        d.paragraph_format.space_after = Pt(4)
        for r in d.runs:
            set_font(r, size=10, color=(70, 70, 70))

    if sub_steps:
        for i, step_text in enumerate(sub_steps, start=1):
            sp = doc.add_paragraph()
            sp.paragraph_format.left_indent = Cm(1.5)
            sp.paragraph_format.first_line_indent = Cm(-0.4)
            sp.paragraph_format.space_before = Pt(1)
            sp.paragraph_format.space_after = Pt(2)
            num_run = sp.add_run(f'{i}.  ')
            set_font(num_run, bold=True, size=10, color=(31, 73, 125))
            text_run = sp.add_run(step_text)
            set_font(text_run, size=10)

    if image_path and os.path.exists(image_path):
        img_para = doc.add_paragraph()
        img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img_para.paragraph_format.left_indent = Cm(1.2)
        img_para.paragraph_format.space_before = Pt(6)
        img_para.paragraph_format.space_after = Pt(2)
        try:
            run = img_para.add_run()
            run.add_picture(image_path, width=Inches(IMG_WIDTH_IN))
        except Exception as e:
            print(f"  Nie udało się wstawić obrazka: {e}")

    if image_caption:
        cap = doc.add_paragraph(f'  {image_caption}')
        cap.paragraph_format.left_indent = Cm(1.2)
        cap.paragraph_format.space_before = Pt(2)
        cap.paragraph_format.space_after = Pt(6)
        for r in cap.runs:
            set_font(r, size=9, italic=True, color=(100, 100, 100))


def add_prompt_block(doc, text):
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Cm(0.5)
    para.paragraph_format.right_indent = Cm(0.5)
    para.paragraph_format.space_before = Pt(8)
    para.paragraph_format.space_after = Pt(8)
    pPr = para._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'EFF3FB')
    pPr.append(shd)
    pBdr = OxmlElement('w:pBdr')
    for side in ['top', 'bottom', 'left', 'right']:
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:space'), '4')
        el.set(qn('w:color'), '4472C4')
        pBdr.append(el)
    pPr.append(pBdr)
    run = para.add_run(text)
    set_font(run, name='Courier New', size=8)


# ── Treść: kroki ─────────────────────────────────────────────────────────────

def build_steps():
    print("\nPrzygotowuję obrazki...")

    img_windsurf_ide = create_windsurf_windows_mockup()

    img_npm = create_powershell_mockup([
        'npm install -g @anthropic-ai/claude-code',
        '>>> added 127 packages in 45s',
        '>>> (instalacja zakończona, może chwilę potrwać)',
    ])

    img_claude = create_powershell_mockup([
        'claude',
        '>>> Opening browser for login...',
        '>>> (przeglądarka się otworzy — zaloguj się)',
    ])

    steps = [
        {
            'title': 'Zainstaluj Windsurf',
            'description': (
                'Windsurf to program, w którym będziesz pracować razem z asystentem AI. '
                'Instaluje się jak każdy zwykły program na Windowsie.'
            ),
            'sub_steps': [
                'Otwórz przeglądarkę (Edge, Chrome — cokolwiek masz).',
                'W pasku adresu u góry wpisz: windsurf.com  i naciśnij Enter.',
                'Na stronie zobaczysz przycisk do pobierania — kliknij go '
                '(szukaj napisu "Download for Windows").',
                'Poczekaj aż plik się pobierze — zobaczysz go na dole przeglądarki '
                'lub w folderze Pobrane (Downloads).',
                'Kliknij dwukrotnie pobrany plik instalatora (.exe).',
                'Jeśli Windows zapyta "czy chcesz zezwolić tej aplikacji..." — kliknij Tak.',
                'Klikaj "Next" przez kolejne ekrany i na końcu "Install".',
                'Po instalacji otwórz Windsurf: kliknij ikonę na Pulpicie '
                'lub wyszukaj "Windsurf" w menu Start (ikona Windows na dole po lewej).',
            ],
            'image_path': None,
            'image_caption': None,
        },
        {
            'title': 'Zainstaluj Node.js',
            'description': (
                'Node.js to silnik, bez którego Claude nie uruchomi się. '
                'Instaluje się jak każdy normalny program — kilka kliknięć.'
            ),
            'sub_steps': [
                'Otwórz przeglądarkę.',
                'Wejdź na stronę: nodejs.org',
                'Zobaczysz dwa przyciski. Kliknij ten zielony — '
                'ma napis "Download Node.js (LTS)" i podpis "Recommended for most users". '
                'NIE klikaj drugiego (Current).',
                'Poczekaj aż plik się pobierze (.msi).',
                'Kliknij dwukrotnie pobrany plik.',
                'Jeśli Windows zapyta "czy chcesz zezwolić..." — kliknij Tak.',
                'Klikaj "Next" przez każdy ekran — nic nie zmieniaj w ustawieniach.',
                'Na końcu kliknij "Install". Może potrwać chwilę.',
                'Gotowe! Możesz zamknąć instalator.',
                'Ważne: uruchom Windsurf ponownie po tej instalacji '
                '(zamknij i otwórz go jeszcze raz), żeby zobaczył nowy Node.js.',
            ],
            'image_path': None,
            'image_caption': None,
        },
        {
            'title': 'Otwórz terminal w Windsurf i wpisz jedną komendę',
            'description': (
                'Terminal to ciemne okienko, w którym wpisujesz polecenia. '
                'W Windsurf jest ukryte na dole ekranu — pokażemy Ci jak je otworzyć.'
            ),
            'sub_steps': [
                'Uruchom Windsurf (jeśli jeszcze nie jest otwarty).',
                'W górnym menu kliknij "Terminal" a potem "New Terminal". '
                'Ewentualnie naciśnij klawisze Ctrl + ` '
                '(tylni apostrof — klawisz pod Esc lub obok cyfry 1, zależy od klawiatury).',
                'Na dole ekranu pojawi się ciemne okienko z niebieskim tekstem — to PowerShell (terminal).',
                'Kliknij w to ciemne okienko żeby je zaznaczyć.',
                'Wpisz poniższą komendę i naciśnij Enter. '
                'Poczekaj — może trwać 1-3 minuty (coś się będzie instalować).',
                'Gdy skończy, w terminalu pojawi się nowa linia z kursorem — to znak że gotowe.',
            ],
            'image_path': img_windsurf_ide,
            'image_caption': (
                'Windsurf na Windows — terminal otworzysz przez menu Terminal → New Terminal. '
                'Zobaczysz niebieski tekst PS C:\\...'
            ),
        },
        {
            'title': 'Otwórz folder Onboarding w Windsurf',
            'description': (
                'Musisz wskazać Windsurfowi folder z plikami onboardingowymi '
                '— ten sam folder, z którego czytasz ten dokument.'
            ),
            'sub_steps': [
                'W Windsurf kliknij "File" w górnym menu.',
                'Kliknij "Open Folder..." (Otwórz folder).',
                'Pojawi się okienko do wyboru folderu — znajdź i zaznacz folder "Onboarding".',
                'Kliknij przycisk "Select Folder" (lub "Otwórz") na dole okienka.',
                'Po lewej stronie Windsurf pojawi się lista plików z folderu Onboarding — to znak że zadziałało.',
            ],
            'image_path': None,
            'image_caption': None,
        },
        {
            'title': 'Uruchom Claude i zaloguj się',
            'description': (
                'Teraz uruchamiamy samego Claude — asystenta AI. '
                'Przy pierwszym uruchomieniu otworzy się przeglądarka żebyś się zalogował.'
            ),
            'sub_steps': [
                'W ciemnym okienku (terminal) na dole Windsurf kliknij żeby je zaznaczyć.',
                'Wpisz: claude  i naciśnij Enter.',
                'Claude może zapytać o zgodę — jeśli zobaczysz pytanie, wpisz Y i Enter.',
                'Automatycznie otworzy się przeglądarka z prośbą o zalogowanie.',
                'Zaloguj się przez swoje konto Anthropic (lub utwórz konto jeśli nie masz).',
                'Wróć do Windsurf — w terminalu zobaczysz że Claude czeka na Twój tekst. '
                'Zobaczysz coś w stylu: "Claude Code >"',
            ],
            'image_path': img_claude,
            'image_caption': 'Tak wygląda uruchomiony Claude w terminalu Windsurf (PowerShell).',
        },
        {
            'title': 'Wklej prompt z Części 2 — Claude zajmie się resztą!',
            'description': (
                'To ostatni krok który robisz sam. '
                'Dajesz Claude instrukcje i od tej chwili on prowadzi Cię przez całą konfigurację '
                '— po polsku, krok po kroku, czekając aż każdy krok zrobisz.'
            ),
            'sub_steps': [
                'Przejdź na następną stronę tego dokumentu.',
                'Zaznacz CAŁY tekst z niebieskiej ramki '
                '(kliknij na początku tekstu, potem Shift+kliknij na końcu ramki).',
                'Skopiuj: Ctrl+C (lub kliknij prawym przyciskiem → Kopiuj).',
                'Wróć do Windsurf i kliknij w ciemne okienko terminala.',
                'Wklej: Ctrl+V (lub kliknij prawym → Wklej).',
                'Naciśnij Enter.',
                'Claude zacznie pisać do Ciebie i przeprowadzi Cię przez każdy następny krok. '
                'Wystarczy że będziesz odpowiadać "gotowe" gdy skończysz każdy krok.',
            ],
            'image_path': None,
            'image_caption': None,
        },
    ]

    # Dodaj komendę npm do Kroku 3 jako osobny obrazek
    steps[2]['image_path_npm'] = img_npm

    return steps


# ── Prompt handoff (Windows) ──────────────────────────────────────────────────

HANDOFF_PROMPT = (
    "Przeprowadź mnie przez konfigurację środowiska Claude Code.\n"
    "\n"
    "Jestem nowym użytkownikiem na WINDOWS 10/11. Właśnie zainstalowałem Windsurf\n"
    "i Claude Code CLI, jestem zalogowany. Pracuję w folderze Onboarding — zawiera\n"
    "podfolder setup/ z gotowymi plikami do skopiowania.\n"
    "\n"
    "Zasady dla każdego kroku:\n"
    "- Tłumacz prosto, bez żargonu technicznego — jakbyś mówił do osoby która\n"
    "  nigdy nie programowała\n"
    "- Używaj komend dla Windows (PowerShell), NIE dla macOS/Linux\n"
    "- Podawaj JEDNĄ komendę lub akcję na raz\n"
    "- Czekaj aż powiem 'gotowe' zanim przejdziesz do następnego kroku\n"
    "- Po każdym kroku sprawdź czy zadziałało\n"
    "- Jeśli coś się nie uda — pomóż mi to naprawić zanim pójdziemy dalej\n"
    "\n"
    "CZĘŚĆ A — Instalacja brakujących narzędzi\n"
    "\n"
    "1. Sprawdź czy mam winget: winget --version\n"
    "   Jeśli nie ma — powiedz żebym zainstalował 'App Installer' z Microsoft Store\n"
    "   (winget jest wbudowany w Windows 10/11, może wymagać aktualizacji)\n"
    "2. Zainstaluj GitHub CLI i jq:\n"
    "   winget install GitHub.cli\n"
    "   winget install jqlang.jq\n"
    "   (gh = do łączenia z GitHub, jq = pomocnik konfiguracyjny)\n"
    "3. Zamknij i otwórz nowy terminal w Windsurf, żeby narzędzia były widoczne\n"
    "4. Zaloguj mnie do GitHub: gh auth login\n"
    "   Wybierz: GitHub.com -> HTTPS -> zaloguj przez przeglądarkę\n"
    "\n"
    "CZĘŚĆ B — Konfiguracja Claude\n"
    "\n"
    "5. Utwórz foldery konfiguracyjne:\n"
    "   $HOME/.claude/skills/   $HOME/.claude/hooks/   $HOME/.claude/memory/\n"
    "6. Skopiuj skille z setup/skills/ do $HOME/.claude/skills/\n"
    "   (state-start, catchup, setup-new-project)\n"
    "7. Skopiuj hook bezpieczeństwa:\n"
    "   setup/hooks/security-scan.py -> $HOME/.claude/hooks/\n"
    "8. Skopiuj status bar:\n"
    "   setup/statusline-command.sh -> $HOME/.claude/\n"
    "9. Wygeneruj plik konfiguracyjny $HOME/.claude/settings.json\n"
    "   z ustawieniami hooka i status baru\n"
    "   (na Windows użyj $env:USERPROFILE przy ścieżkach w PowerShell)\n"
    "10. Zapytaj mnie o:\n"
    "    - w jakim języku chcę rozmawiać z AI (np. polski)\n"
    "    - mój GitHub username\n"
    "    - gdzie trzymam foldery projektów\n"
    "    Skopiuj setup/global-claude-template.md -> $HOME/.claude/CLAUDE.md\n"
    "    i uzupełnij moje dane w miejscach oznaczonych [[...]]\n"
    "11. W pliku $HOME/.claude/skills/setup-new-project/SKILL.md\n"
    "    zamień GITHUB_USERNAME na mój GitHub username\n"
    "12. Weryfikacja: sprawdź że wszystkie pliki istnieją,\n"
    "    pokaż zawartość settings.json, przetestuj /state-start\n"
    "\n"
    "Na końcu wytłumacz jak stworzyć pierwszy projekt:\n"
    "- jak skopiować folder Onboarding jako szablon\n"
    "- jak uruchomić /setup-new-project żeby stworzyć repo na GitHub\n"
    "\n"
    "Pełny opis każdego kroku: setup/onboarding-prompt.md"
)


# ── Generowanie dokumentu ─────────────────────────────────────────────────────

def generate():
    script_dir = os.path.dirname(os.path.abspath(__file__))
    project_root = os.path.dirname(os.path.dirname(script_dir))
    output_path = os.path.join(project_root, 'docs', 'output', 'onboarding-windows.docx')

    print("Generuję onboarding-windows.docx...")
    steps = build_steps()

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    # ── Tytuł ──
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_para.paragraph_format.space_after = Pt(4)
    r = title_para.add_run('Onboarding — Claude Code w Windsurf')
    set_font(r, bold=True, size=20, color=(31, 73, 125))

    sub_para = doc.add_paragraph()
    sub_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_para.paragraph_format.space_after = Pt(14)
    r = sub_para.add_run('Instrukcja startowa · Windows 10 / 11')
    set_font(r, size=11, color=(128, 128, 128))

    intro = doc.add_paragraph(
        'Masz przed sobą 6 prostych kroków. Nie musisz wiedzieć nic o programowaniu — '
        'wystarczy klikać i wpisywać dokładnie to co widzisz. '
        'Po 6. kroku Claude (asystent AI) przejmuje i prowadzi Cię przez resztę po polsku.'
    )
    intro.paragraph_format.space_after = Pt(6)
    for r in intro.runs:
        set_font(r, size=10)

    # ── Część 1 ──
    add_section_heading(doc, 'Część 1 — Zrób to sam  (6 kroków)')

    for i, step in enumerate(steps, start=1):
        add_step(
            doc,
            number=i,
            title=step['title'],
            description=step['description'],
            sub_steps=step.get('sub_steps'),
            image_path=step.get('image_path'),
            image_caption=step.get('image_caption'),
        )
        # Krok 3: dodaj też mockup komendy npm
        if i == 3 and step.get('image_path_npm'):
            npm_para = doc.add_paragraph()
            npm_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            npm_para.paragraph_format.left_indent = Cm(1.2)
            npm_para.paragraph_format.space_before = Pt(4)
            npm_para.paragraph_format.space_after = Pt(2)
            try:
                run = npm_para.add_run()
                run.add_picture(step['image_path_npm'], width=Inches(IMG_WIDTH_IN))
            except Exception:
                pass
            cap = doc.add_paragraph(
                '  Wpisz tę komendę w terminalu i naciśnij Enter.'
            )
            cap.paragraph_format.left_indent = Cm(1.2)
            cap.paragraph_format.space_before = Pt(2)
            cap.paragraph_format.space_after = Pt(6)
            for r in cap.runs:
                set_font(r, size=9, italic=True, color=(100, 100, 100))

    # ── Część 2: Prompt ──
    doc.add_page_break()
    add_section_heading(doc, 'Część 2 — Tekst do wklejenia do Claude')

    instruction = doc.add_paragraph(
        'Zaznacz cały tekst w niebieskiej ramce poniżej, skopiuj (Ctrl+C) '
        'i wklej do ciemnego okienka w Windsurf gdzie działa Claude (Ctrl+V). '
        'Naciśnij Enter. Claude zacznie rozmawiać z Tobą i przeprowadzi Cię przez całą resztę.'
    )
    for r in instruction.runs:
        set_font(r, size=10)

    doc.add_paragraph()
    add_prompt_block(doc, HANDOFF_PROMPT)

    note = doc.add_paragraph(
        'Szczegółowy opis każdego kroku konfiguracji: setup/onboarding-prompt.md'
    )
    note.paragraph_format.space_before = Pt(6)
    for r in note.runs:
        set_font(r, size=9, color=(128, 128, 128))

    # ── Zapis ──
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f'\nZapisano: {output_path}')


if __name__ == '__main__':
    generate()
